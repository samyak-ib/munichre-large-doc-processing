"""The shareable results workbook.

A route's `telemetry.xlsx` accumulates every run made on it, which is the right
shape for comparing experiments and the wrong shape for showing someone where
the project stands. This writes one self-contained workbook for one batch: what
it scored, what the score rests on, and what it cost.

The Assumptions sheet is the part that matters. Every equivalence the scorer
honours is a judgement about golden data rather than about the extraction, so
each one is listed with the accuracy it is worth — measured, not asserted, by
scoring the same rows again with that rule switched off.
"""

from __future__ import annotations

import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook

from .accuracy import ACCURACY_COLUMNS, DEFAULT_POLICY, POLICY_FLAGS, load_golden
from .report import _write_sheet, _write_summary_block
from .rescore import RunRecord, policy_influence
from .schema_loader import TableSchema
from .telemetry import CALL_COLUMNS

INFLUENCE_COLUMNS = (
    "assumption",
    "flag",
    "why",
    "rows_golden",
    "rows_matched_with",
    "rows_matched_without",
    "cells_scored_with",
    "cells_scored_without",
    "accuracy_with_pct",
    "accuracy_without_pct",
    "influence_pp",
)

DOCUMENT_COLUMNS = (
    "document",
    "pages",
    "model",
    "rows_golden",
    "rows_matched",
    "row_recall_pct",
    "cells_scored",
    "cells_correct",
    "cell_accuracy_pct",
    "exact_row_pct",
)

# Why each assumption exists, in the words a first-time reader needs. Keyed by
# the ScoringPolicy flag it belongs to.
ASSUMPTION_NOTES = {
    "ignore_golden_blank": (
        "Golden leaves a cell empty where the document does carry a value. Whether "
        "our answer is right is unknowable from this data, so the cell is counted "
        "and reported but never scored — it can neither help nor hurt."
    ),
    "money_empty_is_zero": (
        "The schema tells the model to return N/A when it finds no amount; golden "
        "writes 0. Scoring that disagreement measures the two conventions rather "
        "than the extraction."
    ),
    "leading_zero_tolerant_ids": (
        "Excel stored numeric claim ids as numbers, so golden holds 40512146091 "
        "where the document prints 040512146091. Without this the row does not "
        "match at all, which is why switching it off can raise cell accuracy "
        "while matching fewer rows."
    ),
    "status_synonyms": (
        "Golden records claim status as-is, so CLOSED, C and 'Settled / Closed' "
        "all appear. They are one status."
    ),
    "state_synonyms": "TX and Texas are the same state.",
    "semantic_description": (
        "Golden clips the description column at 50 characters and varies in case "
        "and punctuation. Two descriptions that state the same thing count as "
        "equal; a narrative and a taxonomy code do not."
    ),
}

# Judgements that shaped the extraction rather than the scoring. They carry no
# measurable delta — they are here because a reader deserves to know the number
# rests on them.
STANDING_CAVEATS = (
    (
        "Occurrence ID is derived from the claim number",
        "The documents scored here carry no occurrence column. Golden derives the "
        "occurrence from the claim number by dropping its sequence suffix "
        "(C00320453-02 -> C00320453), and the extraction prompt now says so. That "
        "rule was learned from this golden set and may not hold on a document "
        "that prints a real occurrence column.",
    ),
    (
        "Description means the coded cause, not the narrative",
        "Where a document carries both a coded cause column (STRAIN OR INJURY BY "
        "TWISTING) and a free-text narrative (WHILE WORKING ON HIS KNEES...), "
        "golden holds the coded cause. schema.json lists 'Accident Description' as "
        "an allowed source, which points at the narrative. This is a recommended "
        "correction to the class definition, not a silent fix.",
    ),
    (
        "Insurer Loss Run is the largest remaining error, for two separate reasons",
        "It is not fixed here and is the obvious next candidate. In one document "
        "the carrier name appears only in a letterhead image — 'FCCI' occurs zero "
        "times in the text layer — so a model that reads the page as an image "
        "gets it and a model that reads the text layer returns N/A. In another, "
        "the text carries both the brand (CHUBB) and the legal entity (FEDERAL "
        "INSURANCE COMPANY); golden holds the brand and both models returned the "
        "entity. Neither is a transcription failure.",
    ),
    (
        "QA can only fix what the text layer can confirm",
        "The review proposes corrections; one is written into the table only when "
        "the value it proposes occurs in the document's own text layer. That is "
        "what stops a confident-sounding invention from reaching the deliverable, "
        "and it has a cost: on a scanned document there is no text layer, so every "
        "finding is reported and none is applied. The Issues sheet's "
        "`qa_unverified` rows are where that shows up.",
    ),
    (
        "Golden is the arbiter, and it is a small sample",
        "Five documents and 92 claim rows. Every rule above was written after "
        "reading these documents, so the reported accuracy is an upper bound on "
        "what an unseen loss run would score.",
    ),
    (
        "Cost is a lower bound where usage is missing",
        "A provider that returns no token counts leaves those calls priced at "
        "zero. The Telemetry sheet's calls_without_usage column says how many.",
    ),
)


@dataclass
class BatchResults:
    """Everything one batch of documents produced, ready to write."""

    batch_id: str
    records: list[RunRecord]
    schema: TableSchema
    golden_path: Path
    ledger_rows: Sequence[dict[str, Any]] = ()
    call_rows: Sequence[Sequence[Any]] = ()

    def document_rows(self) -> list[dict[str, Any]]:
        rows = []
        for record in self.records:
            golden = load_golden(self.golden_path, record.document)
            if not golden:
                continue
            for model, result in record.score_all(golden, self.schema, DEFAULT_POLICY).items():
                rows.append(
                    {
                        "document": record.document,
                        "pages": record.pages,
                        "model": model,
                        "rows_golden": result.rows_golden,
                        "rows_matched": result.rows_matched,
                        "row_recall_pct": round(result.row_recall, 1),
                        "cells_scored": result.cells_compared,
                        "cells_correct": result.cells_correct,
                        "cell_accuracy_pct": round(result.cell_accuracy, 1),
                        "exact_row_pct": round(result.exact_row_rate, 1),
                    }
                )
        return rows

    def column_rows(self) -> list[dict[str, Any]]:
        """Per-column accuracy, pooled across models and documents.

        Pooled rather than per-run: the question this sheet answers is which
        columns the extractor is bad at, and one document's 2 rows should not
        read as loudly as another's 51.
        """
        totals: dict[str, list[int]] = {}
        for record in self.records:
            golden = load_golden(self.golden_path, record.document)
            if not golden:
                continue
            for result in record.score_all(golden, self.schema, DEFAULT_POLICY).values():
                for name, entry in result.columns.items():
                    bucket = totals.setdefault(name, [0, 0])
                    bucket[0] += entry.compared
                    bucket[1] += entry.correct
        rows = [
            {
                "column": name,
                "compared": compared,
                "correct": correct,
                "accuracy_pct": round(correct / compared * 100, 1) if compared else 0.0,
            }
            for name, (compared, correct) in totals.items()
        ]
        return sorted(rows, key=lambda r: r["accuracy_pct"])

    def headline(self) -> dict[str, Any]:
        rows = self.document_rows()
        cells = sum(r["cells_scored"] for r in rows)
        correct = sum(r["cells_correct"] for r in rows)
        golden = sum(r["rows_golden"] for r in rows)
        matched = sum(r["rows_matched"] for r in rows)
        return {
            "batch_id": self.batch_id,
            "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "documents": len({r["document"] for r in rows}),
            "models": ", ".join(sorted({r["model"] for r in rows})),
            "golden_rows": golden,
            "rows_matched": matched,
            "row_recall_pct": round(matched / golden * 100, 1) if golden else 0.0,
            "cells_scored": cells,
            "cells_correct": correct,
            "cell_accuracy_pct": round(correct / cells * 100, 1) if cells else 0.0,
        }


def batch_filename(version: str = "v1", stamp: str | None = None) -> str:
    """The workbook's name, fixed once per batch so it can be rewritten in place."""
    return f"{version}_{stamp or time.strftime('%Y%m%d_%H%M%S')}.xlsx"


def write_results(
    results: BatchResults,
    out_dir: Path,
    version: str = "v1",
    filename: str | None = None,
) -> Path:
    """Write `<out_dir>/<version>_<date>_<time>.xlsx` and return its path.

    Passing `filename` rewrites an existing workbook rather than minting a new
    one, which is what lets a batch publish after every document instead of only
    at the end.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = workbook.create_sheet("Summary")
    _write_summary_block(summary, results.headline())
    _write_sheet(
        summary,
        list(DOCUMENT_COLUMNS),
        [[row.get(c, "") for c in DOCUMENT_COLUMNS] for row in results.document_rows()],
        start_row=len(results.headline()) + 3,
    )

    _write_sheet(
        workbook.create_sheet("Accuracy By Column"),
        ["column", "compared", "correct", "accuracy_pct"],
        [
            [row["column"], row["compared"], row["correct"], row["accuracy_pct"]]
            for row in results.column_rows()
        ],
    )

    _write_assumptions(workbook.create_sheet("Assumptions"), results)

    telemetry = workbook.create_sheet("Telemetry")
    _write_sheet(
        telemetry,
        list(ACCURACY_COLUMNS),
        [[row.get(c, "") for c in ACCURACY_COLUMNS] for row in results.ledger_rows],
    )
    _write_sheet(
        telemetry,
        list(CALL_COLUMNS),
        [list(row) for row in results.call_rows],
        start_row=len(results.ledger_rows) + 3,
    )

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / (filename or batch_filename(version))
    workbook.save(path)
    return path


def _write_assumptions(sheet, results: BatchResults) -> None:
    influence = policy_influence(
        results.records, results.golden_path, results.schema, POLICY_FLAGS
    )
    for row in influence:
        row["why"] = ASSUMPTION_NOTES.get(row["flag"], "")

    _write_sheet(
        sheet,
        list(INFLUENCE_COLUMNS),
        [[row.get(c, "") for c in INFLUENCE_COLUMNS] for row in influence],
    )
    _write_sheet(
        sheet,
        ["standing caveat", "detail"],
        [list(pair) for pair in STANDING_CAVEATS],
        start_row=len(influence) + 3,
    )
