"""Reading a batch's telemetry back out of the cumulative ledger."""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor

from openpyxl import Workbook, load_workbook

from lossrun.report import append_ledger, read_batch_telemetry, write_workbook
from lossrun.telemetry import CALL_COLUMNS, SUMMARY_COLUMNS


def build_ledger(path, runs: list[dict], calls: list[dict]) -> None:
    book = Workbook()
    book.remove(book.active)
    runs_sheet = book.create_sheet("Runs")
    runs_sheet.append(list(SUMMARY_COLUMNS))
    for row in runs:
        runs_sheet.append([row.get(c, "") for c in SUMMARY_COLUMNS])
    accuracy_sheet = book.create_sheet("Accuracy")
    accuracy_sheet.append(["batch_id", "run_id", "document"])
    calls_sheet = book.create_sheet("Calls")
    calls_sheet.append(list(CALL_COLUMNS))
    for row in calls:
        calls_sheet.append([row.get(c, "") for c in CALL_COLUMNS])
    book.save(path)


def test_calls_are_returned_for_a_document_with_no_golden_entry(tmp_path):
    """A document with no golden entry never appears in the `Accuracy` sheet at
    all, but it still made real API calls that cost real money — those must
    not be silently dropped from the batch's telemetry."""
    path = tmp_path / "telemetry.xlsx"
    build_ledger(
        path,
        runs=[{"batch_id": "b1", "run_id": "r1", "document": "no-golden.pdf"}],
        calls=[{"run_id": "r1", "document": "no-golden.pdf", "stage": "extract", "latency_s": 5.0}],
    )

    accuracy, calls, runs = read_batch_telemetry(path, "b1")

    assert accuracy == []
    assert len(runs) == 1
    assert len(calls) == 1
    document_index = CALL_COLUMNS.index("document")
    assert calls[0][document_index] == "no-golden.pdf"


def test_calls_outside_the_wanted_batch_are_excluded(tmp_path):
    path = tmp_path / "telemetry.xlsx"
    build_ledger(
        path,
        runs=[
            {"batch_id": "b1", "run_id": "r1", "document": "a.pdf"},
            {"batch_id": "b2", "run_id": "r2", "document": "b.pdf"},
        ],
        calls=[
            {"run_id": "r1", "document": "a.pdf", "stage": "extract", "latency_s": 1.0},
            {"run_id": "r2", "document": "b.pdf", "stage": "extract", "latency_s": 1.0},
        ],
    )

    _, calls, runs = read_batch_telemetry(path, "b1")

    assert len(runs) == 1
    assert len(calls) == 1
    document_index = CALL_COLUMNS.index("document")
    assert calls[0][document_index] == "a.pdf"


def test_no_matching_batch_returns_empty_everything(tmp_path):
    path = tmp_path / "telemetry.xlsx"
    build_ledger(path, runs=[{"batch_id": "b1", "run_id": "r1", "document": "a.pdf"}], calls=[])

    accuracy, calls, runs = read_batch_telemetry(path, "unknown-batch")

    assert (accuracy, calls, runs) == ([], [], [])


def test_a_missing_ledger_file_returns_empty_everything(tmp_path):
    accuracy, calls, runs = read_batch_telemetry(tmp_path / "nope.xlsx", "b1")
    assert (accuracy, calls, runs) == ([], [], [])


def test_append_ledger_is_safe_under_concurrent_writers(tmp_path):
    """A parallel batch calls `append_ledger` once per document, concurrently,
    against the same file. Without a lock this races: an interleaved
    load-modify-save loses updates or corrupts the workbook."""
    path = tmp_path / "telemetry.xlsx"
    n = 12

    with ThreadPoolExecutor(max_workers=n) as pool:
        list(
            pool.map(
                lambda i: append_ledger(
                    path, {"batch_id": "b1", "run_id": f"r{i}", "document": f"doc{i}.pdf"}, ()
                ),
                range(n),
            )
        )

    runs_sheet = load_workbook(path)["Runs"]
    assert runs_sheet.max_row == n + 1  # header row + one per call, none lost
    documents = {row[SUMMARY_COLUMNS.index("document")].value for row in runs_sheet.iter_rows(min_row=2)}
    assert documents == {f"doc{i}.pdf" for i in range(n)}


def test_control_characters_are_stripped_instead_of_failing_the_write(tmp_path):
    """OCR on a noisy scan can hand the model a NULL byte, a lone UTF-16
    surrogate, or a Unicode noncharacter. openpyxl's own check only catches
    the first kind; its lxml backend rejects all three at save time (`All
    strings must be XML compatible`), which used to fail the whole document
    after a correct, already-paid-for extraction. Stripping all three is
    better than losing the run."""
    path = tmp_path / "extraction.xlsx"
    dirty = "Cash\x00App claim\x0bnote" + chr(0xD800) + "X" + chr(0xFFFE) + chr(0xFFFF)

    write_workbook(
        path,
        final_rows=[{"Claimant Name": dirty}],
        final_columns=["Claimant Name"],
        raw_rows=[],
        raw_columns=[],
        issues=[],
        calls=[],
        summary={"document": "d.pdf"},
    )

    sheet = load_workbook(path)["Final Table"]
    assert sheet.cell(row=2, column=1).value == "CashApp claimnoteX"
