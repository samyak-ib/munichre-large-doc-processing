"""Scoring an extraction against golden data.

The comparison has to be forgiving about *format* and strict about *value* — a
date written two ways is not an error, and a wrong number is not a formatting
quirk. Getting that backwards silently mis-states the only number that matters.
"""

from __future__ import annotations

from datetime import datetime

import pytest

from lossrun.accuracy import _as_date, _as_money, load_golden, score, values_match
from lossrun.schema_loader import load_schema

SCHEMA = load_schema()


def golden_row(**overrides) -> dict[str, str]:
    row = {
        "Claim Number": "C003",
        "Claimant Name": "McAllister, John",
        "Policy Number": "P-1",
        "Accident Date": "06/28/2022",
        "Indemnity Paid": "1234.56",
        "Claim Status": "CLOSED",
    }
    row.update(overrides)
    return row


def test_dates_match_across_the_formats_each_side_uses():
    assert values_match("Accident Date", "Jan 5, 2023", "01/05/2023")
    assert values_match("Accident Date", "2023-01-05 00:00:00", "01/05/2023")
    assert values_match("Accident Date", "01/05/2023", "01/05/2023")
    assert not values_match("Accident Date", "01/06/2023", "01/05/2023")


def test_month_name_dates_parse_whole_not_first_token():
    # "Jan 5, 2023".split(" ")[0] is "Jan" — parsing that loses the date.
    assert _as_date("Jan 5, 2023") == datetime(2023, 1, 5).date()
    assert _as_date("January 5, 2023") == datetime(2023, 1, 5).date()
    assert _as_date("2023-01-05 00:00:00") == datetime(2023, 1, 5).date()


def test_money_matches_across_currency_decoration():
    assert values_match("Indemnity Paid", "$1,234.56", "1234.56")
    assert values_match("Indemnity Paid", "1234.560", "1234.56")
    assert values_match("Indemnity Paid", "(500.00)", "-500")
    assert not values_match("Indemnity Paid", "1234.56", "975.39")


def test_money_parsing_handles_blanks_and_junk():
    assert _as_money("$0.00") == 0.0
    assert _as_money("") is None
    assert _as_money("see note") is None


def test_text_matches_ignoring_case_and_spacing():
    assert values_match("Claim Status", "closed", "CLOSED")
    assert values_match("Claimant Name", "McAllister,  John", "McAllister, John")
    assert not values_match("Claimant Name", "MacAllister, John", "McAllister, John")


def test_a_cell_empty_on_both_sides_is_not_scored():
    result = score([golden_row(Description="N/A")], [golden_row(Description="N/A")], SCHEMA, "m")
    assert "Description" not in {m["column"] for m in result.mismatches}
    assert result.columns["Description"].compared == 0


def test_a_cell_golden_leaves_blank_is_reported_but_not_scored():
    """Golden is incomplete in places, so a blank there proves nothing.

    Whether our value is right is unknowable from this data, so the cell is
    counted separately and left out of the score entirely.
    """
    extracted = [golden_row(**{"Loss State": "FL"})]
    golden = [golden_row(**{"Loss State": "N/A"})]
    result = score(extracted, golden, SCHEMA, "m")
    assert result.columns["Loss State"].compared == 0
    assert result.spurious == 1
    assert result.cell_accuracy == 100.0, "an unscorable cell cannot drag the score down"


def test_a_blank_money_cell_and_an_explicit_zero_agree():
    """The schema says return N/A when nothing is found; golden writes 0.

    Scoring that disagreement measures the two conventions, not the extraction.
    """
    assert values_match("Expense Reserved", "N/A", "0")
    assert values_match("Expense Reserved", "0", "N/A")
    assert values_match("Expense Reserved", "", "0.00")
    assert values_match("Recovery Deductible", "$0.00", "N/A")
    assert values_match("Indemnity Paid", "0", "0.0")


def test_empty_versus_zero_leniency_does_not_extend_to_real_amounts():
    assert not values_match("Indemnity Paid", "N/A", "1234.56")
    assert not values_match("Indemnity Paid", "1234.56", "N/A")
    assert not values_match("Indemnity Paid", "0", "1234.56")


def test_empty_versus_zero_leniency_is_money_only():
    # "0" is not a claim number, a status, or a date.
    assert not values_match("Claim Number", "N/A", "0")
    assert not values_match("Claim Status", "", "0")


def test_a_zero_for_blank_row_can_now_be_fully_correct():
    extracted = [golden_row(**{"Expense Reserved": "N/A", "Recovery Deductible": "N/A"})]
    golden = [golden_row(**{"Expense Reserved": "0", "Recovery Deductible": "0"})]
    result = score(extracted, golden, SCHEMA, "m")
    assert result.cell_accuracy == 100.0
    assert result.exact_row_rate == 100.0


def test_a_perfect_extraction_scores_100():
    rows = [golden_row()]
    result = score(rows, [dict(rows[0])], SCHEMA, "m")
    assert result.rows_matched == 1
    assert result.cell_accuracy == 100.0
    assert result.exact_row_rate == 100.0
    assert result.mismatches == []


def test_rows_are_matched_on_claim_number_not_the_merge_key():
    """A wrong claimant name must not hide the row from scoring — that is
    precisely the error the score needs to report."""
    extracted = [golden_row(**{"Claimant Name": "Robyn Streight"})]
    result = score(extracted, [golden_row()], SCHEMA, "m")
    assert result.rows_matched == 1
    assert result.columns["Claimant Name"].correct == 0
    assert result.exact_rows == 0


def test_missing_and_extra_rows_are_counted_separately():
    golden = [golden_row(), golden_row(**{"Claim Number": "C004"})]
    extracted = [golden_row(), golden_row(**{"Claim Number": "C999"})]
    result = score(extracted, golden, SCHEMA, "m")
    assert result.rows_matched == 1
    assert result.missing_keys == ["c004"]
    assert result.extra_keys == ["c999"]
    assert result.row_recall == 50.0
    assert result.row_precision == 50.0


def test_golden_lookup_tolerates_a_renamed_document(tmp_path):
    from openpyxl import Workbook

    from lossrun.accuracy import GOLDEN_TO_SCHEMA

    path = tmp_path / "gt.xlsx"
    book = Workbook()
    sheet = book.active
    header = ["Filename"] + list(GOLDEN_TO_SCHEMA)
    sheet.append(header)
    sheet.append(["guid__10306283_Updated Acords LRs_Application_CAU CPP.PDF"] + ["x"] * len(GOLDEN_TO_SCHEMA))
    book.save(path)

    assert load_golden(path, "LRs_Application_CAU CPP.PDF") is not None
    assert load_golden(path, "SomeOtherDocument.pdf") is None


def test_claim_status_codes_and_words_are_equivalent():
    """Golden records status as-is: CLOSED, C, Closed and "C (Closed)" all occur.

    A model writing one form where golden holds another is right, not wrong.
    """
    for form in ("CLOSED", "C", "Closed", "C (Closed)", "Settled / Closed"):
        assert values_match("Claim Status", form, "CLOSED"), form
    assert values_match("Claim Status", "O", "OPEN")
    assert not values_match("Claim Status", "OPEN", "CLOSED")


def test_state_names_and_codes_are_equivalent():
    assert values_match("Loss State", "Texas", "TX")
    assert values_match("Loss State", "TX", "Texas")
    assert values_match("Loss State", "north carolina", "NC")
    assert not values_match("Loss State", "TX", "NC")


def test_an_unrecognised_status_falls_back_to_text_comparison():
    assert values_match("Claim Status", "Litigation", "litigation")
    assert not values_match("Claim Status", "Litigation", "CLOSED")


def test_a_claim_id_excel_coerced_to_a_number_still_matches():
    """Golden stores 040512146091 as the number 40512146091.

    The document prints the leading zero and the model transcribes it, so a
    literal comparison would score a correct extraction as zero.
    """
    extracted = [golden_row(**{"Claim Number": "040512146091"})]
    golden = [golden_row(**{"Claim Number": "40512146091"})]
    result = score(extracted, golden, SCHEMA, "m")
    assert result.rows_matched == 1
    assert result.missing_keys == []
    assert result.extra_keys == []
    assert result.columns["Claim Number"].correct == 1


def test_genuinely_different_claim_ids_do_not_match():
    extracted = [golden_row(**{"Claim Number": "40512146099"})]
    golden = [golden_row(**{"Claim Number": "40512146091"})]
    result = score(extracted, golden, SCHEMA, "m")
    assert result.rows_matched == 0
    assert result.missing_keys and result.extra_keys


def test_leading_zero_tolerance_is_limited_to_numeric_ids():
    # Alphanumeric ids keep strict comparison: 0ABC and ABC are different claims.
    assert not values_match("Claim Number", "0ABC123", "ABC123")


def test_accuracy_rows_carry_the_page_count():
    """A score is only interpretable against document size.

    A 2-page document and a 38-page scan sit in the same sheet; without pages
    the reader cannot tell which is which.
    """
    from lossrun.accuracy import ACCURACY_COLUMNS

    assert "pages" in ACCURACY_COLUMNS
    result = score([golden_row()], [golden_row()], SCHEMA, "m")
    row = result.summary_row(document="d.pdf", pages=38)
    assert row["pages"] == 38
    assert set(ACCURACY_COLUMNS) - set(row) <= {"batch_id", "run_id", "started_at", "effort"}


def test_a_wrapped_identifier_scores_as_the_identifier_it_is():
    """A PDF wraps a long policy number mid-string, so the text layer reads
    `001- WC19A-78355` where the table means `001-WC19A-78355`. The row key
    already treats those as one claim; the cell comparison must agree, or a
    correct transcription is scored wrong inside the row it matched."""
    assert values_match("Policy Number", "001- WC19A-78355", "001-WC19A-78355")
    assert values_match("Claim Number", "C00320453 -02", "C00320453-02")
    assert not values_match("Policy Number", "001-WC19A-78355", "001-WC20A-78355")


# --- matching a document to its golden entry ---------------------------------


def _resolve(candidates, document):
    from lossrun.accuracy import _resolve_golden_filename

    return _resolve_golden_filename(candidates, document)


def test_an_exact_filename_always_wins_over_a_containment_match():
    """The failure this prevents is quiet and total.

    With 29 documents, names nest: `Loss Runs.pdf` is a substring of `WC Loss
    Runs.pdf`, `GLI Loss Runs.PDF` and six more. A containment-first rule scored
    it against 155 golden rows drawn from eight different documents.
    """
    candidates = [
        "Loss Runs.pdf",
        "WC Loss Runs.pdf",
        "GLI Loss Runs.PDF",
        "Auto Loss Runs.pdf",
        "2017-22 CIC Pkg Loss Runs.PDF",
    ]
    assert _resolve(candidates, "Loss Runs.pdf") == "Loss Runs.pdf"
    assert _resolve(candidates, "WC Loss Runs.pdf") == "WC Loss Runs.pdf"
    assert _resolve(candidates, "GLI Loss Runs.PDF") == "GLI Loss Runs.PDF"


def test_a_renamed_document_still_finds_its_golden_entry():
    """The case the loose match exists for: golden keeps the original name."""
    candidates = ["093fca2c-8c12__Updated Acords LRs_Application_CAU CPP.PDF", "Loss Runs.pdf"]
    assert _resolve(candidates, "LRs_Application_CAU CPP.PDF") == candidates[0]


def test_the_entry_carrying_the_whole_document_name_wins():
    """Two golden entries transcribe the same PDF and they disagree.

    One is the GUID-prefixed original containing the document name in full; the
    other is a shorter name merely contained in it. The first is the specific
    match, and picking between them cannot be left to sheet order.
    """
    guid = "093fca2c__Updated Acords LRs_Application_CAU CPP MAR PKG WCO Loss Runs.PDF"
    candidates = [guid, "CAU CPP MAR PKG WCO Loss Runs.PDF", "Loss Runs.pdf"]
    assert _resolve(candidates, "LRs_Application_CAU CPP MAR PKG WCO Loss Runs.PDF") == guid
    # …and that shorter entry is still reachable by its own exact name.
    assert _resolve(candidates, "CAU CPP MAR PKG WCO Loss Runs.PDF") == candidates[1]


def test_a_document_with_no_golden_entry_resolves_to_nothing():
    assert _resolve(["Loss Runs.pdf"], "Loss-3.pdf") is None
    assert _resolve([], "anything.pdf") is None


def test_every_sheet_of_the_golden_workbook_is_read(tmp_path):
    """The set grew a second sheet; reading only the first scores the new
    documents as having no golden at all rather than failing loudly."""
    from openpyxl import Workbook

    from lossrun.accuracy import load_golden

    path = tmp_path / "golden.xlsx"
    book = Workbook()
    book.remove(book.active)
    header = ["Filename", "Claim ID", "Claimant Name", "Loss State"]
    first = book.create_sheet("Sheet1")
    first.append(header)
    first.append(["old.pdf", "C1", "A", "TX"])
    second = book.create_sheet("Sheet2")
    second.append(header)
    second.append(["new.pdf", "C2", "B", "CA"])
    book.save(path)

    assert len(load_golden(path, "old.pdf")) == 1
    assert len(load_golden(path, "new.pdf")) == 1, "the second sheet must be read"
    assert load_golden(path, "absent.pdf") is None


# --- row-level accuracy -------------------------------------------------------


def test_row_accuracy_weights_every_row_equally():
    """A row with 20 scorable cells and one with 2 count the same.

    Pooled cell accuracy answers "how many cells are right"; this answers "how
    correct is a typical row", which is the number a reviewer feels.
    """
    from lossrun.accuracy import score
    from lossrun.schema_loader import load_schema

    schema = load_schema()
    golden = [
        # Four scorable cells, all of them right.
        {"Claim Number": "C1", "Claimant Name": "A", "Loss State": "TX", "Policy Number": "P"},
        # Three scorable cells — golden is blank for Policy Number, so it is not
        # scored — and only the claim number is right.
        {"Claim Number": "C2", "Claimant Name": "B", "Loss State": "CA", "Policy Number": "N/A"},
    ]
    extracted = [
        dict(golden[0]),
        {**golden[1], "Claimant Name": "WRONG", "Loss State": "NY"},
    ]
    result = score(extracted, golden, schema, "m")

    assert result.rows_matched == 2
    assert result.row_accuracies == pytest.approx([100.0, 33.3], abs=0.1)
    # Pooled cells lean toward the row carrying more of them...
    assert result.cell_accuracy == pytest.approx(71.4, abs=0.1)  # 5 of 7
    # ...while the row-level figure gives each row one vote.
    assert result.row_accuracy == pytest.approx(66.7, abs=0.1)
    assert result.row_accuracy_overall == pytest.approx(66.7, abs=0.1)


def test_a_row_we_never_found_scores_zero_in_the_overall_figure_only():
    from lossrun.accuracy import score
    from lossrun.schema_loader import load_schema

    schema = load_schema()
    golden = [
        {"Claim Number": "C1", "Claimant Name": "A", "Loss State": "TX"},
        {"Claim Number": "C2", "Claimant Name": "B", "Loss State": "CA"},
    ]
    result = score([dict(golden[0])], golden, schema, "m")
    assert result.rows_matched == 1
    assert result.row_accuracy == 100.0, "of what we returned, all of it is right"
    assert result.row_accuracy_overall == 50.0, "half the document never came back"
