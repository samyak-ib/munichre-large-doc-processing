"""End-to-end pipeline wiring, with the API stubbed out.

Exercises ingest -> profile -> chunk -> extract -> merge -> qa -> verify -> report
against a synthetic loss run, so the plumbing is checked without spending tokens.
"""

from __future__ import annotations

import json
import re

import fitz
import pytest
from openpyxl import load_workbook

from lossrun import pipeline
from lossrun.config import ApiConfig, ChunkingConfig, Config, Pricing, ReasoningConfig
from lossrun.superapp_client import RunResult

CLAIMS = [
    ("P-100", "C003", "McAllister, John"),
    ("P-100", "C004", "Okonkwo, Ada"),
    ("P-100", "C005", "Rivera-Sánchez, Luis"),
    ("P-200", "C006", "Brightwater Holdings LLC"),
]


def make_config(tmp_path, *, qa_enabled: bool = True, **chunk_overrides) -> Config:
    chunking = ChunkingConfig(
        max_pages_per_chunk=2,
        overlap_pages=1,
        header_pages=1,
        overlap_anchor_rows=5,
        max_request_bytes=20 * 1024 * 1024,
        max_resumes=2,
        **chunk_overrides,
    )
    return Config(
        api=ApiConfig(
            base_url="https://example.invalid/api/v1",
            poll_interval_s=0,
            poll_deadline_s=10,
            request_timeout_s=10,
            max_concurrent_calls=2,
            max_retries=0,
        ),
        primary_model="model-a",
        chunking=chunking,
        model_overrides={},
        reasoning=ReasoningConfig(),
        pricing={"model-a": Pricing(input=0.20, output=0.80)},
        token="test-token",
        qa_enabled=qa_enabled,
    )


def write_loss_run(path, pages: int = 4) -> None:
    """A multi-page PDF whose text layer contains the claim rows."""
    doc = fitz.open()
    per_page = max(1, len(CLAIMS) // pages + 1)
    for index in range(pages):
        page = doc.new_page()
        lines = [
            "ACME LOGISTICS INC - LOSS RUN",
            "Valuation Date: 01/15/2026",
            "Policy Number | Claim Number | Claimant | Paid",
        ]
        for policy, claim, name in CLAIMS[index * per_page : (index + 1) * per_page]:
            lines.append(f"{policy} | {claim} | {name} | $1,200.00")
        page.insert_text((50, 60), "\n".join(lines), fontsize=9)
    doc.save(path)
    doc.close()


class FakeClient:
    """Stands in for SuperAppClient, returning canned layout and table payloads."""

    def __init__(self, *, config, telemetry, behaviour=None):
        self.config = config
        self.telemetry = telemetry
        self.behaviour = behaviour or {}
        self.calls: list[dict] = []
        self._resumed = False

    def __enter__(self):
        return self

    def __exit__(self, *_exc):
        return None

    def close(self):
        pass

    def run(
        self, *, model, prompt, instructions="", attachments=None, stage="", label="",
        pages="", effort=None,
    ):
        self.calls.append(
            {
                "model": model,
                "stage": stage,
                "label": label,
                "prompt": prompt,
                "effort": effort,
                "pages": pages,
                "attachments": len(attachments or []),
            }
        )
        text = self._body(stage, prompt, model)
        result = RunResult(
            response_id=f"resp_{len(self.calls)}",
            status="completed",
            output_text=text,
            input_tokens=1000,
            output_tokens=200,
            poll_count=1,
            latency_s=0.5,
        )
        self.telemetry.record(
            stage=stage,
            model=model,
            label=label,
            pages=pages,
            attempt=1,
            effort=effort or "",
            response_id=result.response_id,
            status=result.status,
            latency_s=result.latency_s,
            poll_count=result.poll_count,
            input_tokens=result.input_tokens,
            output_tokens=result.output_tokens,
        )
        return result

    def _body(self, stage, prompt, model):
        if stage == "layout":
            return json.dumps(
                {
                    "present_columns": {"Claim Number": "Claim Number"},
                    "policy_number_placement": "column",
                    "date_format": "MM/DD/YYYY",
                    "document_values": {
                        "Insured": "ACME LOGISTICS INC",
                        "Valuation Date": "01/15/2026",
                    },
                }
            )
        if stage == "qa":
            return json.dumps(
                self.behaviour.get("qa_payload") or {"findings": [], "missing_rows": []}
            )
        if self.behaviour.get("truncate_first") and not self._resumed:
            self._resumed = True
            return self._table(CLAIMS[:2], truncated=True)
        if self.behaviour.get("hallucinates"):
            mutated = list(CLAIMS[:-1]) + [("P-200", "C0000006", "Brightwater Holdings LLC")]
            return self._table(mutated)
        if self.behaviour.get("misreads_a_claimant"):
            return self._table([("P-100", "C003", "MacAllister, John")] + list(CLAIMS[1:]))
        if self.behaviour.get("rows_per_chunk"):
            # Split the table across chunks so each merged row is attributed to
            # the chunk it was first read from.
            match = re.search(r"chunk (\d+) of", prompt)
            index = int(match.group(1)) if match else 1
            return self._table(CLAIMS[:2] if index == 1 else CLAIMS[2:])
        return self._table(CLAIMS)

    @staticmethod
    def _table(claims, truncated: bool = False) -> str:
        columns = ["Policy Number", "Claim Number", "Claimant Name", "Indemnity Paid"]
        rows = [[p, c, n, "$1,200.00"] for p, c, n in claims]
        payload = {
            "columns": columns,
            "rows": rows,
            "truncated": truncated,
            "last_row_key": list(claims[-1]) if claims else [],
        }
        return json.dumps(payload)


@pytest.fixture
def stub_client(monkeypatch):
    created: list[FakeClient] = []

    def install(**behaviour):
        def factory(*, config, telemetry):
            client = FakeClient(config=config, telemetry=telemetry, behaviour=behaviour)
            created.append(client)
            return client

        monkeypatch.setattr(pipeline, "ModelRouter", factory)
        return created

    return install


def test_single_shot_document_produces_the_full_table(tmp_path, stub_client):
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.route == "single_shot"
    assert outcome.rows == len(CLAIMS)
    assert outcome.unverified == 0
    assert outcome.cost_usd > 0
    assert outcome.workbook.exists()
    assert outcome.ledger.exists()


def test_a_run_is_filed_under_its_route_with_that_routes_ledger(tmp_path, stub_client):
    """Each route keeps its own runs and its own ledger, and nothing else.

    Runs measured through SuperApp carry agent-loop tokens a direct call never
    pays for, so one ledger holding both would invite an invalid comparison.
    """
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)
    out = tmp_path / "out"

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=out, log=lambda *_: None
    )

    assert outcome.run_dir.parent == out / "superapp_calls", "the stub routes to superapp"
    assert outcome.ledger == out / "superapp_calls" / "telemetry.xlsx"
    assert not (out / "telemetry_ledger.xlsx").exists(), "no ledger above the route folders"

    from openpyxl import load_workbook

    assert load_workbook(outcome.ledger)["Runs"].max_row == 2, "one header plus one run"


def test_chunked_document_deduplicates_the_seam(tmp_path, stub_client):
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=5)  # 5 pages at 2 per chunk -> multiple overlapping chunks

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.route == "chunked"
    assert outcome.chunks > 1
    # Every chunk returns the same rows; merging on the key collapses them to one set.
    assert outcome.rows == len(CLAIMS)


def test_truncated_chunk_is_resumed(tmp_path, stub_client):
    created = stub_client(truncate_first=True)
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf,
        config=make_config(tmp_path),
        out_dir=tmp_path / "out",
        log=lambda *_: None,
    )

    assert outcome.rows == len(CLAIMS)
    labels = [c["label"] for c in created[0].calls]
    assert any("resume" in label for label in labels)
    resume_prompt = next(c["prompt"] for c in created[0].calls if "resume" in c["label"])
    assert "RESUMING" in resume_prompt


def test_workbook_carries_every_sheet_and_column(tmp_path, stub_client):
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    book = load_workbook(outcome.workbook)
    assert book.sheetnames == ["Final Table", "Raw Rows", "Issues", "Telemetry"]

    sheet = book["Final Table"]
    header = [c.value for c in sheet[1]]
    assert len(header) == 25
    assert sheet.max_row == len(CLAIMS) + 1

    rows = {r[header.index("Claim Number")] for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows == {claim for _, claim, _ in CLAIMS}
    # Document-level values are stamped onto every row.
    insured = {r[header.index("Insured")] for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert insured == {"ACME LOGISTICS INC"}


def test_a_hallucinated_key_is_caught_against_the_text_layer(tmp_path, stub_client):
    stub_client(hallucinates=True)
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.unverified == 1
    assert "key_not_found" in _issue_categories(outcome.workbook)


def test_ledger_appends_one_row_per_run(tmp_path, stub_client):
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)
    out_dir = tmp_path / "out"

    for _ in range(2):
        outcome = pipeline.run_document(
            pdf, config=make_config(tmp_path), out_dir=out_dir, log=lambda *_: None
        )

    book = load_workbook(outcome.ledger)
    assert book["Runs"].max_row == 3  # header + two runs
    assert book["Calls"].max_row > 3


def _issue_categories(workbook_path) -> set[str]:
    sheet = load_workbook(workbook_path)["Issues"]
    header = [c.value for c in sheet[1]]
    index = header.index("category")
    return {row[index] for row in sheet.iter_rows(min_row=2, values_only=True)}


def test_a_chunk_that_parses_to_nothing_saves_its_raw_response(tmp_path, stub_client, monkeypatch):
    """An unparseable chunk must leave the model's text behind.

    Re-running one to find out what it said costs a full chunk of tokens, so the
    evidence has to survive the run that produced it.
    """
    stub_client()
    monkeypatch.setattr(
        FakeClient, "_body", lambda self, stage, prompt, model: (
            json.dumps({"present_columns": {}}) if stage == "layout" else "I cannot find a table."
        )
    )
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf,
        config=make_config(tmp_path),
        out_dir=tmp_path / "out",
        log=lambda *_: None,
    )

    assert outcome.rows == 0
    saved = list((outcome.run_dir / "raw").glob("raw-*.txt"))
    assert saved, "expected the raw response to be saved"
    assert saved[0].read_text() == "I cannot find a table."


def test_a_clean_run_leaves_no_raw_debris(tmp_path, stub_client):
    stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)
    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None,
    )
    assert outcome.rows > 0
    assert not (outcome.run_dir / "raw").exists()


def test_ledger_migrates_when_a_column_is_added(tmp_path):
    """A ledger written before a column existed must not shift its old rows.

    Appending the new row shape under an old header silently misaligns every
    value past the insertion point — corrupting the cross-run comparison the
    ledger exists for.
    """
    from openpyxl import Workbook, load_workbook

    from lossrun.report import append_ledger
    from lossrun.telemetry import SUMMARY_COLUMNS

    ledger = tmp_path / "telemetry_ledger.xlsx"
    old_columns = [c for c in SUMMARY_COLUMNS if c not in ("effort", "calls_without_usage")]

    book = Workbook()
    book.remove(book.active)
    runs = book.create_sheet("Runs")
    runs.append(old_columns)
    runs.append(["old-run" if c == "run_id" else 7 if c == "rows" else "" for c in old_columns])
    calls = book.create_sheet("Calls")
    calls.append(["run_id"])
    book.save(ledger)

    append_ledger(ledger, {"run_id": "new-run", "rows": 51, "effort": "low"}, [])

    sheet = load_workbook(ledger)["Runs"]
    header = [c.value for c in sheet[1]]
    assert header == list(SUMMARY_COLUMNS)

    rows = {r[header.index("run_id")]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["old-run"][header.index("rows")] == 7, "the old row kept its values"
    # openpyxl reads a written empty string back as None; either means "blank".
    assert not rows["old-run"][header.index("effort")], "the new column is blank, not shifted"
    assert rows["new-run"][header.index("effort")] == "low"
    assert rows["new-run"][header.index("rows")] == 51


def test_an_empty_response_is_retried_not_accepted_as_zero_rows(tmp_path, stub_client, monkeypatch):
    """A silent agent route returns completed with no text.

    Taking that as "this document has no claims" loses every row after paying
    for the upload, so the chunk is retried instead.
    """
    created = stub_client()
    state = {"calls": 0}

    def body(self, stage, prompt, model):
        if stage == "layout":
            return json.dumps({"present_columns": {}})
        if stage == "qa":
            return json.dumps({"findings": [], "missing_rows": []})
        state["calls"] += 1
        return "" if state["calls"] == 1 else FakeClient._table(CLAIMS)

    monkeypatch.setattr(FakeClient, "_body", body)
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf,
        config=make_config(tmp_path),
        out_dir=tmp_path / "out",
        log=lambda *_: None,
    )

    assert outcome.rows == len(CLAIMS), "the retry recovered the chunk"
    assert state["calls"] == 2


# --- QA review ---------------------------------------------------------------


def _qa_finding(row, column, correct_value, reason="checked against the page"):
    return {
        "findings": [
            {"row": list(row), "column": column, "correct_value": correct_value, "reason": reason}
        ],
        "missing_rows": [],
    }


def _final_table(workbook_path) -> list[dict]:
    sheet = load_workbook(workbook_path)["Final Table"]
    header = [c.value for c in sheet[1]]
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_a_correction_the_document_supports_is_applied(tmp_path, stub_client):
    """The reviewer's whole purpose: a misread name the page can settle."""
    stub_client(
        misreads_a_claimant=True,
        qa_payload=_qa_finding(
            ("P-100", "C003", "MacAllister, John"), "Claimant Name", "McAllister, John"
        ),
    )
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.qa_findings == 1
    assert outcome.qa_applied == 1
    names = {row["Claimant Name"] for row in _final_table(outcome.workbook)}
    assert "McAllister, John" in names
    assert "MacAllister, John" not in names
    assert "qa_corrected" in _issue_categories(outcome.workbook)


def test_a_correction_the_document_does_not_support_is_reported_not_applied(tmp_path, stub_client):
    """The guard that replaces adjudication's never-invent-a-value property.

    QA can propose any string it likes, so a proposal that appears nowhere in the
    document must not reach the table — it is a hallucination with a confident
    tone, and the extracted value has at least been read off a page.
    """
    stub_client(
        misreads_a_claimant=True,
        qa_payload=_qa_finding(
            ("P-100", "C003", "MacAllister, John"), "Claimant Name", "Nobody, Invented"
        ),
    )
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.qa_findings == 1
    assert outcome.qa_applied == 0
    names = {row["Claimant Name"] for row in _final_table(outcome.workbook)}
    assert "Nobody, Invented" not in names
    assert "MacAllister, John" in names, "the extracted value stands"
    assert "qa_unverified" in _issue_categories(outcome.workbook)


def test_a_proposal_to_clear_a_cell_is_reported_not_applied(tmp_path, stub_client):
    """Deleting a value is held to the same proof as replacing one.

    Measured over the five golden documents, unproven clearings were the only
    thing this stage got wrong — see docs/CHALLENGES.md #22.
    """
    stub_client(
        qa_payload=_qa_finding(
            ("P-100", "C003", "McAllister, John"), "Indemnity Paid", "N/A"
        )
    )
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.qa_findings == 1
    assert outcome.qa_applied == 0
    row = next(r for r in _final_table(outcome.workbook) if r["Claim Number"] == "C003")
    assert row["Indemnity Paid"] == "1200"
    assert "qa_unverified" in _issue_categories(outcome.workbook)


def test_a_row_the_review_found_on_the_page_is_flagged_not_added(tmp_path, stub_client):
    stub_client(
        qa_payload={
            "findings": [],
            "missing_rows": [{"row": ["P-300", "C009", "Unseen Claimant"], "reason": "on page 2"}],
        }
    )
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    assert outcome.rows == len(CLAIMS), "a reported row is never added to the table"
    assert "qa_row_missing" in _issue_categories(outcome.workbook)


def test_each_chunks_rows_are_reviewed_against_that_chunks_pages(tmp_path, stub_client):
    """A review only means something if it is looking at the right pages."""
    created = stub_client(rows_per_chunk=True)
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=5)

    pipeline.run_document(
        pdf, config=make_config(tmp_path), out_dir=tmp_path / "out", log=lambda *_: None
    )

    qa_calls = [c for c in created[0].calls if c["stage"] == "qa"]
    assert len(qa_calls) == 2, "one call per chunk that produced rows, not per chunk"
    assert all(c["attachments"] == 1 for c in qa_calls), "the pages are re-attached"
    # The rows read from the first chunk are reviewed against the first chunk.
    first = next(c for c in qa_calls if c["pages"] == "1-2")
    assert "C003" in first["prompt"] and "C005" not in first["prompt"]


def test_qa_makes_no_calls_when_it_is_off(tmp_path, stub_client):
    created = stub_client()
    pdf = tmp_path / "loss_run.pdf"
    write_loss_run(pdf, pages=2)

    outcome = pipeline.run_document(
        pdf,
        config=make_config(tmp_path, qa_enabled=False),
        out_dir=tmp_path / "out",
        log=lambda *_: None,
    )

    assert not [c for c in created[0].calls if c["stage"] == "qa"]
    assert outcome.qa_findings == 0
    assert outcome.rows == len(CLAIMS)
